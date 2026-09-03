"""DOCX parser (python-docx).

Reads paragraphs (preserving heading levels) and tables. Images in DOCX are
ignored in the MVP but the registry structure supports them.
"""
from __future__ import annotations

import io

from app.parsers.base import DocumentParser, ParseError
from app.schemas.domain import (
    ContentType,
    ParsedDocument,
    Region,
    RegionType,
    SourceType,
    TableStructure,
)

_HEADING_STYLES = {
    "title": RegionType.TITLE,
    "heading 1": RegionType.HEADING,
    "heading 2": RegionType.HEADING,
    "heading 3": RegionType.HEADING,
    "heading 4": RegionType.HEADING,
    "heading 5": RegionType.HEADING,
    "heading 6": RegionType.HEADING,
}


class DocxParser(DocumentParser):
    source_type = SourceType.DOCX

    async def parse(self, data: bytes) -> ParsedDocument:
        try:
            import docx
            doc = docx.Document(io.BytesIO(data))
        except Exception as exc:
            raise ParseError(f"Could not parse DOCX: {exc}") from exc

        document_id = "doc_placeholder"
        regions: list[Region] = []
        section_stack: list[str] = []
        order = 0

        def to_hierarchy() -> dict[str, str]:
            return {f"h{i+1}": lvl for i, lvl in enumerate(section_stack)}

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name or "").lower() if para.style else ""
            base = style.split(" ")[0]
            if style.startswith("heading") or style == "title":
                level = style.replace("heading ", "")
                try:
                    level = int(level) if level.isdigit() else 1
                except Exception:
                    level = 1
                rtype = _HEADING_STYLES.get(style, RegionType.HEADING)
                while len(section_stack) >= max(level, 1):
                    section_stack.pop()
                section_stack.append(text)
                regions.append(
                    Region(document_id=document_id, page_number=1, type=rtype, content=text,
                           reading_order=order, hierarchy=to_hierarchy(),
                           metadata={"style": style}, source_type=SourceType.DOCX)
                )
            else:
                rtype = RegionType.PARAGRAPH
                regions.append(
                    Region(document_id=document_id, page_number=1, type=rtype, content=text,
                           reading_order=order, hierarchy=to_hierarchy(), source_type=SourceType.DOCX)
                )
            order += 1

        for table in doc.tables:
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            header = rows[0]
            data_rows = rows[1:]
            summary = f"table: columns {header} rows {data_rows}"
            regions.append(
                Region(document_id=document_id, page_number=1, type=RegionType.TABLE,
                       content_type=ContentType.TABLE, content=summary, reading_order=order,
                       table=TableStructure(columns=header, rows=data_rows, summary=summary),
                       source_type=SourceType.DOCX, hierarchy=to_hierarchy())
            )
            order += 1

        return ParsedDocument(document_id=document_id, source_type=SourceType.DOCX,
                              pages=1, regions=regions, metadata={"format": "docx"})
